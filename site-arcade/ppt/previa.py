#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Desenha uma previa de cada slide do PPTX.

O LibreOffice nao converte pptx neste ambiente, entao a conferencia visual e
feita aqui: le o arquivo com python-pptx e redesenha as formas, imagens e
textos na posicao real. As imagens saem identicas (sao as mesmas do deck); o
texto e aproximado, em DejaVu, o que basta para achar transbordo, sobreposicao
e elemento fora da margem.

Tambem reporta, em texto, os problemas que dao para medir: forma fora do
slide, margem menor que a minima e texto que nao cabe na caixa.
"""
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

ARQ = sys.argv[1] if len(sys.argv) > 1 else 'portfolio-adilson-serafim.pptx'
ESC = 110          # pixels por polegada na previa
MARGEM_MIN = 0.45  # polegadas
F_REG = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
F_BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
pol = lambda emu: Emu(emu).inches


def fonte(tam_pt, negrito):
    return ImageFont.truetype(F_BOLD if negrito else F_REG, max(7, int(tam_pt * ESC / 72)))


def cor_de(obj, padrao=(255, 255, 255)):
    try:
        rgb = obj.rgb
        return (rgb[0], rgb[1], rgb[2])
    except Exception:
        return padrao


def quebrar(d, texto, f, larg):
    linhas = []
    for bruta in texto.split('\n'):
        atual = ''
        for p in bruta.split(' '):
            t = (atual + ' ' + p).strip()
            if d.textlength(t, font=f) <= larg or not atual:
                atual = t
            else:
                linhas.append(atual); atual = p
        linhas.append(atual)
    return linhas


def main():
    pres = Presentation(ARQ)
    LP, AP = pol(pres.slide_width), pol(pres.slide_height)
    problemas = []

    for n, slide in enumerate(pres.slides, 1):
        im = Image.new('RGB', (int(LP * ESC), int(AP * ESC)), (0, 0, 0))
        d = ImageDraw.Draw(im)
        caixas = []

        for sh in slide.shapes:
            if sh.left is None:
                continue
            x, y = pol(sh.left), pol(sh.top)
            w, h = pol(sh.width), pol(sh.height)

            if x < -0.01 or y < -0.01 or x + w > LP + 0.01 or y + h > AP + 0.01:
                problemas.append('slide %d: forma fora do slide (%.2f,%.2f %.2fx%.2f)'
                                 % (n, x, y, w, h))
            elif min(x, y, LP - (x + w), AP - (y + h)) < MARGEM_MIN - 0.01:
                problemas.append('slide %d: a %.2f" da borda — %s'
                                 % (n, min(x, y, LP - (x + w), AP - (y + h)),
                                    (sh.text_frame.text[:34] if sh.has_text_frame else sh.shape_type)))

            cx = (int(x * ESC), int(y * ESC), int((x + w) * ESC), int((y + h) * ESC))

            if sh.shape_type == 13 or sh.__class__.__name__ == 'Picture':   # imagem
                try:
                    foto = Image.open(__import__('io').BytesIO(sh.image.blob)).convert('RGBA')
                    foto = foto.resize((max(1, cx[2] - cx[0]), max(1, cx[3] - cx[1])))
                    im.paste(foto, (cx[0], cx[1]), foto)
                except Exception as e:
                    d.rectangle(cx, outline=(90, 90, 90))
                continue

            preenche = None
            try:
                if sh.fill.type is not None and sh.fill.type == 1:
                    preenche = cor_de(sh.fill.fore_color)
            except Exception:
                pass
            linha = None
            try:
                if sh.line.fill.type == 1:
                    linha = cor_de(sh.line.color, (120, 120, 120))
            except Exception:
                pass
            if preenche or linha:
                d.rectangle(cx, fill=preenche, outline=linha)

            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text
                p0 = sh.text_frame.paragraphs[0]
                r0 = p0.runs[0] if p0.runs else None
                tam = (r0.font.size.pt if r0 is not None and r0.font.size else 12)
                neg = bool(r0.font.bold) if r0 is not None else False
                c = cor_de(r0.font.color, (255, 255, 255)) if r0 is not None else (255, 255, 255)
                f = fonte(tam, neg)
                linhas = quebrar(d, t, f, (w * ESC) - 2)
                alt_linha = f.size * 1.32
                precisa = len(linhas) * alt_linha / ESC
                if precisa > h + 0.04:
                    problemas.append('slide %d: texto nao cabe (precisa %.2f", caixa %.2f") — "%s"'
                                     % (n, precisa, h, t[:44].replace('\n', ' ')))
                al = (p0.alignment.__str__() if p0.alignment else '').upper()
                yy = cx[1]
                for ln in linhas:
                    lx = cx[0] + 1
                    if 'CENTER' in al:
                        lx = cx[0] + ((w * ESC) - d.textlength(ln, font=f)) / 2
                    elif 'RIGHT' in al:
                        lx = cx[2] - d.textlength(ln, font=f) - 1
                    d.text((lx, yy), ln, font=f, fill=c)
                    yy += alt_linha
                caixas.append((n, cx, t[:30]))

        im.save('previa-%02d.png' % n)

    print('previas geradas: %d' % n)
    if problemas:
        print('\nPROBLEMAS (%d):' % len(problemas))
        for p in problemas:
            print(' -', p)
    else:
        print('\nnenhum problema de medida encontrado')


if __name__ == '__main__':
    main()
